#!/usr/bin/env python3
"""Genera js/catalog-data.js y js/products-data.js a partir del Excel de productos.

Uso:  python tools/build_catalog.py [ruta/al/excel.xlsx]
      (por defecto lee tools/productos.xlsx)

Reglas:
  - Solo productos con activo = VERDADERO. Los inactivos no se publican en ningún lado.
  - Se descartan filas sin descripción o sin filtro1 (filas vacías del Excel).
  - filtro1 -> categoría (menú), filtro2 -> subcategoría, filtro3 -> agrupador (chips).
  - filtro4 no se usa (en el Excel tiene espesores / fechas mal convertidas).
  - Largo / peso que Excel convirtió en fecha se muestran vacíos (no se inventa el dato).
"""
import datetime, json, os, re, sys, unicodedata
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "tools", "productos.xlsx")

# ---- Categorías: orden, grupo de menú, slug, nombre para mostrar, portada -----------------
GROUPS = ["Aceros y perfiles", "Chapas y cubiertas", "Construcción y cerramientos", "Ferretería y complementos"]
CATS = [
    # filtro1 en el Excel,            grupo, slug,                 nombre,                                    portada
    ("ANG/PLAN/T/U",                    0, "angulos-planchuelas",  "Ángulos, planchuelas y perfiles T/U",     "barr.png"),
    ("Perfiles Laminados",              0, "perfiles-laminados",   "Perfiles laminados",                      "aceros-p-la-construccion.png"),
    ("Perfiles Conformados",            0, "perfiles-c",           "Perfiles conformados",                    "perdiles-c.png"),
    ("Laminados Red. Y Cuad.",          0, "laminados",            "Laminados redondos y cuadrados",          "barr.png"),
    ("Trefilados",                      0, "trefilados",           "Trefilados",                              "barr.png"),
    ("Caños",                           0, "canos",                "Caños",                                   "canos.png"),
    ("Tubos S/C",                       0, "tubos-sin-costura",    "Tubos sin costura",                       "angulos.png"),
    ("Acero Inoxidable",                0, "acero-inoxidable",     "Acero inoxidable",                        "Aceros.png"),
    ("Aluminios",                       0, "aluminios",            "Aluminios",                               "Aceros.png"),
    ("Chapas Lisas",                    1, "chapas-lisas",         "Chapas lisas",                            "chapas-lisas.png"),
    ("Chapas Conformadas",              1, "chapas-conformadas",   "Chapas conformadas",                      "chapas-techos.png"),
    ("Complementos p/Ch. de Techo",     1, "complementos-techo",   "Complementos para chapa de techo",        "chapas-techos.png"),
    ("Aislantes",                       1, "aislantes",            "Aislantes",                               "aislantes.png"),
    ("Hierros y Mallas Construccion",   2, "hierros-mallas",       "Hierros y mallas para construcción",      "hierros-y-mallas.png"),
    ("Mallas Job y Metal Desplegado",   2, "mallas-job",           "Mallas Job y metal desplegado",           "tejidos.png"),
    ("Tejidos y Telas Metalicas",       2, "tejidos",              "Tejidos y telas metálicas",               "tejidos.png"),
    ("Alambres y Clavos",               2, "alambres-clavos",      "Alambres y clavos",                       "clavos-y-alambres.png"),
    ("Tornillos y Fijaciones",          3, "tornillos",            "Tornillos y fijaciones",                  "tornillos.png"),
    ("Complementos Herreria",           3, "herreria",             "Complementos de herrería",                "tornillos.png"),
    ("Ferreteria",                      3, "ferreteria",           "Ferretería",                              "aceros-p-la-construccion.png"),
]
COVER_DIR = "assets/productos/"

WORDS = {
    "Angulos": "Ángulos", "Conduccion": "Conducción", "Mecanico": "Mecánico", "Metricas": "Métricas",
    "Fundicion": "Fundición", "Zingueria": "Zinguería", "Eolicos": "Eólicos", "Herreria": "Herrería",
    "Ferreteria": "Ferretería", "Perfileria": "Perfilería", "Metalicas": "Metálicas", "Artistico": "Artístico",
    "Plasticas": "Plásticas", "Laminas": "Láminas", "Construccion": "Construcción", "Baston": "Bastón",
}
EXACT = {
    "Plan. SAE 1045": "Planchuelas SAE 1045", "Plan. Perforadas": "Planchuelas perforadas",
    "RED.": "Redondos", "RECT.": "Rectangulares", "CUADR.": "Cuadrados", "1.x50x6": "1.50x6.00 Mts",
    "Ch. Conformada Cincalum": "Chapa conformada cincalum",
    "Ch. Conformada Color": "Chapa conformada color",
    "Ch. Conformada Plasticas": "Chapa conformada plástica",
    "Ch. Lisa Color en Hojas": "Chapa lisa color en hojas",
    "Lam. Frio": "Laminado en frío", "Lam. Caliente": "Laminado en caliente",
    "Acc. p/Cables y Cadenas": "Accesorios para cables y cadenas",
    "Complementos p/Perfileria Alum.": "Complementos para perfilería de aluminio",
    "Complementos p/Aislantes": "Complementos para aislantes",
    "Accesorios p/Alambrados": "Accesorios para alambrados",
    "Accesorios p/Tejidos": "Accesorios para tejidos",
    "Caños Ac. Inoxidable": "Caños de acero inoxidable",
    "Chapas Ac. Inoxidable": "Chapas de acero inoxidable",
    "Ac. Inoxidable Varios": "Acero inoxidable varios",
    "Ch. Antidesliz. de Aluminio": "Chapas antideslizantes de aluminio",
    "Chapas Plasticas/Policarbonato": "Chapas plásticas y policarbonato",
}


def pretty(s):
    if s in EXACT:
        return EXACT[s]
    return " ".join(WORDS.get(w, w) for w in s.split(" "))


def slug(s):
    s = unicodedata.normalize("NFD", s.lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def text(v):
    """Texto si es una etiqueta real; None si está vacío, es número o fecha mal convertida."""
    if v is None or v == "-":
        return None
    if isinstance(v, (datetime.datetime, int, float)):
        return None
    s = str(v).strip()
    if re.fullmatch(r"[0-9.,/\"' ]+(mm)?", s):
        return None
    return s


def num(v):
    if v is None or isinstance(v, datetime.datetime):
        return None
    try:
        n = float(str(v).replace(",", "."))
    except ValueError:
        return None
    return n if n > 0 else None


wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb.active
head = [c.value for c in ws[1]]
ix = {h: i for i, h in enumerate(head)}
rows = [r for r in ws.iter_rows(min_row=2, values_only=True)]

by_f1 = {c[0]: i for i, c in enumerate(CATS)}
stats = dict(total=len(rows), inactivos=0, sin_datos=0, cat_desconocida=set(), largo_roto=0, kg_roto=0)
cats = [{"id": c[2], "label": c[3], "group": c[1], "image": COVER_DIR + c[4], "sub": [], "_subs": {}} for c in CATS]
products = []

for r in rows:
    if r[ix["activo"]] is not True:
        stats["inactivos"] += 1
        continue
    desc = str(r[ix["descripcion"]]).strip() if r[ix["descripcion"]] else ""
    f1 = r[ix["filtro1"]]
    if not desc or f1 in (None, "-"):
        stats["sin_datos"] += 1
        continue
    if f1 not in by_f1:
        stats["cat_desconocida"].add(f1)
        continue
    ci = by_f1[f1]
    cat = cats[ci]
    f2 = text(r[ix["filtro2"]]) or cat["label"]
    f3 = text(r[ix["filtro3"]])
    sub = cat["_subs"].get(f2)
    if sub is None:
        sub = {"label": pretty(f2), "count": 0, "_groups": {}}
        cat["_subs"][f2] = sub
    sub["count"] += 1
    gname = pretty(f3) if f3 else None
    sub["_groups"][gname] = sub["_groups"].get(gname, 0) + 1
    raw_l, raw_k = r[ix["largoStd"]], r[ix["kgUnidad"]]
    largo, kg = num(raw_l), num(raw_k)
    if isinstance(raw_l, datetime.datetime):
        stats["largo_roto"] += 1
    if isinstance(raw_k, datetime.datetime):
        stats["kg_roto"] += 1
    products.append([int(r[ix["id"]]), desc, ci, f2, gname, r[ix["venta_por"]] or "", largo, kg])

# ---- subcategorías: orden por cantidad, ids únicos; agrupadores solo si hay 2 o más --------
for cat in cats:
    subs = sorted(cat["_subs"].items(), key=lambda kv: -kv[1]["count"])
    used = set()
    for raw, s in subs:
        sid = slug(s["label"]) or "sub"
        while sid in used:
            sid += "-2"
        used.add(sid)
        groups = s["_groups"]
        garr = []
        if len(groups) >= 2:
            ordered = sorted(groups.items(), key=lambda kv: (kv[0] is None, -kv[1]))
            garr = [{"id": slug(g) if g else "otros", "label": g or "Otros", "count": n} for g, n in ordered]
        cat["sub"].append({"id": sid, "label": s["label"], "count": s["count"], "groups": garr, "_raw": raw})
    del cat["_subs"]
    cat["count"] = sum(s["count"] for s in cat["sub"])
    names = [s["label"] for s in cat["sub"]][:4]
    cat["intro"] = "Incluye " + ", ".join(n if i == 0 else n[0].lower() + n[1:] for i, n in enumerate(names)) + (" y más." if len(cat["sub"]) > 4 else ".")

# ---- productos compactos: [id, desc, catIdx, subIdx, grpIdx(-1), ventaPor, largo, kg] ------
compact = []
for pid, desc, ci, f2raw, gname, venta, largo, kg in products:
    cat = cats[ci]
    si = next(i for i, s in enumerate(cat["sub"]) if s["_raw"] == f2raw)
    gi = -1
    for j, g in enumerate(cat["sub"][si]["groups"]):
        if g["label"] == (gname or "Otros"):
            gi = j
    compact.append([pid, desc, ci, si, gi, venta, largo, kg])
for cat in cats:
    for s in cat["sub"]:
        del s["_raw"]

# categorías sin productos activos no se publican; reindexar
keep = [i for i, c in enumerate(cats) if c["count"] > 0]
remap = {old: new for new, old in enumerate(keep)}
cats = [cats[i] for i in keep]
for p in compact:
    p[2] = remap[p[2]]


def dump(o):
    return json.dumps(o, ensure_ascii=False, separators=(",", ":"))


catalog_js = (
    "/* Domissi Hermanos — categorías del catálogo.\n"
    "   ARCHIVO GENERADO por tools/build_catalog.py a partir del Excel de productos. No editar a mano. */\n"
    "window.DOMISSI_GROUPS = " + dump(GROUPS) + ";\n"
    "window.DOMISSI_CATALOG = " + dump(cats) + ";\n"
)
products_js = (
    "/* Domissi Hermanos — productos activos.\n"
    "   ARCHIVO GENERADO por tools/build_catalog.py. Fila: [código, descripción, categoría, subcategoría, agrupador, venta por, largo (m), kg/un.] */\n"
    "window.DOMISSI_PRODUCTS = " + dump(compact) + ";\n"
)
with open(os.path.join(ROOT, "js", "catalog-data.js"), "w", encoding="utf-8") as f:
    f.write(catalog_js)
with open(os.path.join(ROOT, "js", "products-data.js"), "w", encoding="utf-8") as f:
    f.write(products_js)

print("filas Excel:", stats["total"], "| inactivas:", stats["inactivos"], "| vacías/sin categoría:", stats["sin_datos"])
print("publicados:", len(compact), "| categorías:", len(cats), "| largo perdido (fecha):", stats["largo_roto"], "| kg perdido (fecha):", stats["kg_roto"])
if stats["cat_desconocida"]:
    print("CATEGORIAS SIN MAPEAR!", stats["cat_desconocida"])
for c in cats:
    print(f'{c["count"]:4d}  {c["label"]}  ->  ' + " | ".join(
        f'{s["label"]}({s["count"]})' + (f'[{len(s["groups"])}g]' if s["groups"] else "") for s in c["sub"]))
